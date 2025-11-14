import requests
import json
import base64
from datetime import datetime
import os
import csv
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import tkinter as tk
from tkinter import messagebox, filedialog
import threading

class GeradorRelatorios:
    def __init__(self, api_base='http://localhost:5000/api', username='admin', password='password'):
        self.API_BASE = api_base
        self.USERNAME = username
        self.PASSWORD = password
        self.headers = {
            'Content-Type': 'application/json',
            'Authorization': 'Basic ' + base64.b64encode(f'{self.USERNAME}:{self.PASSWORD}'.encode()).decode()
        }
        
        # Configurar fontes para PDF
        try:
            pdfmetrics.registerFont(TTFont('Arial', 'arial.ttf'))
            pdfmetrics.registerFont(TTFont('Arial-Bold', 'arialbd.ttf'))
        except:
            print("Fontes Arial não encontradas, usando fontes padrão")

    def testar_conexao_api(self):
        """Testa a conexão com a API"""
        try:
            response = requests.get(f"{self.API_BASE}/health", timeout=5)
            return response.status_code == 200
        except:
            return False

    def obter_dados_api(self, endpoint):
        """Obtém dados de um endpoint específico da API"""
        try:
            response = requests.get(f"{self.API_BASE}/{endpoint}", headers=self.headers, timeout=10)
            if response.status_code == 200:
                return response.json()
            else:
                print(f"Erro na API: {response.status_code}")
                return None
        except Exception as e:
            print(f"Erro ao obter dados: {e}")
            return None

    def obter_estatisticas(self):
        """Obtém estatísticas do sistema"""
        try:
            response = requests.get(f"{self.API_BASE}/info", headers=self.headers, timeout=10)
            if response.status_code == 200:
                return response.json()
            return None
        except:
            return None

    def gerar_relatorio_completo(self, formato='pdf', salvar_como=None):
        """Gera relatório completo do sistema"""
        if not self.testar_conexao_api():
            raise Exception("Não foi possível conectar à API. Verifique se o servidor está rodando.")

        # Coletar todos os dados
        estatisticas = self.obter_estatisticas()
        caixas = self.obter_dados_api('caixas') or []
        arquivos = self.obter_dados_api('arquivos') or []
        segurados = self.obter_dados_api('segurados') or []

        if formato.lower() == 'pdf':
            return self._gerar_pdf_completo(estatisticas, caixas, arquivos, segurados, salvar_como)
        elif formato.lower() == 'excel':
            return self._gerar_excel_completo(estatisticas, caixas, arquivos, segurados, salvar_como)
        else:
            raise ValueError("Formato deve ser 'pdf' ou 'excel'")

    def gerar_relatorio_caixas(self, formato='pdf', salvar_como=None):
        """Gera relatório específico de caixas"""
        if not self.testar_conexao_api():
            raise Exception("Não foi possível conectar à API.")

        caixas = self.obter_dados_api('caixas') or []
        estatisticas = self.obter_estatisticas()

        if formato.lower() == 'pdf':
            return self._gerar_pdf_caixas(caixas, estatisticas, salvar_como)
        elif formato.lower() == 'excel':
            return self._gerar_excel_caixas(caixas, estatisticas, salvar_como)
        else:
            raise ValueError("Formato deve ser 'pdf' ou 'excel'")

    def gerar_relatorio_arquivos(self, formato='pdf', salvar_como=None):
        """Gera relatório específico de arquivos"""
        if not self.testar_conexao_api():
            raise Exception("Não foi possível conectar à API.")

        arquivos = self.obter_dados_api('arquivos') or []
        estatisticas = self.obter_estatisticas()

        if formato.lower() == 'pdf':
            return self._gerar_pdf_arquivos(arquivos, estatisticas, salvar_como)
        elif formato.lower() == 'excel':
            return self._gerar_excel_arquivos(arquivos, estatisticas, salvar_como)
        else:
            raise ValueError("Formato deve ser 'pdf' ou 'excel'")

    def gerar_relatorio_segurados(self, formato='pdf', salvar_como=None):
        """Gera relatório específico de segurados"""
        if not self.testar_conexao_api():
            raise Exception("Não foi possível conectar à API.")

        segurados = self.obter_dados_api('segurados') or []
        estatisticas = self.obter_estatisticas()

        if formato.lower() == 'pdf':
            return self._gerar_pdf_segurados(segurados, estatisticas, salvar_como)
        elif formato.lower() == 'excel':
            return self._gerar_excel_segurados(segurados, estatisticas, salvar_como)
        else:
            raise ValueError("Formato deve ser 'pdf' ou 'excel'")

    def _gerar_pdf_completo(self, estatisticas, caixas, arquivos, segurados, salvar_como):
        """Gera PDF do relatório completo"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Completo_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

        doc = SimpleDocTemplate(salvar_como, pagesize=A4)
        elements = []

        # Estilos
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            alignment=1,  # Centro
            textColor=colors.HexColor('#2c3e50')
        )

        # Título
        title = Paragraph("RELATÓRIO COMPLETO - SISTEMA GCA", title_style)
        elements.append(title)

        # Data de geração
        data_geracao = Paragraph(f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", styles['Normal'])
        elements.append(data_geracao)
        elements.append(Spacer(1, 20))

        # Estatísticas
        elements.append(Paragraph("ESTATÍSTICAS DO SISTEMA", styles['Heading2']))
        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            stats_data = [
                ['📦 Caixas Cadastradas', stats.get('caixas', 0)],
                ['📁 Arquivos Cadastrados', stats.get('arquivos', 0)],
                ['👤 Segurados Cadastrados', stats.get('segurados', 0)],
                ['🕐 Uptime do Sistema', estatisticas.get('uptime', 'N/A')]
            ]
            stats_table = Table(stats_data, colWidths=[300, 100])
            stats_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#ecf0f1')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 10),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(stats_table)
        elements.append(Spacer(1, 20))

        # Caixas
        elements.append(Paragraph("CAIXAS CADASTRADAS", styles['Heading2']))
        if caixas:
            caixas_data = [['Código', 'NBI', 'NBF', 'Prateleira', 'Bloco', 'Andar', 'Corredor']]
            for caixa in caixas:
                caixas_data.append([
                    caixa.get('codigo', ''),
                    caixa.get('NBI', ''),
                    caixa.get('NBF', ''),
                    caixa.get('prateleira', ''),
                    caixa.get('bloco', ''),
                    caixa.get('andar', ''),
                    caixa.get('corredor', '')
                ])
            caixas_table = Table(caixas_data, repeatRows=1)
            caixas_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e74c3c')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f9ebea')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(caixas_table)
        else:
            elements.append(Paragraph("Nenhuma caixa cadastrada.", styles['Normal']))
        elements.append(Spacer(1, 20))

        # Arquivos
        elements.append(Paragraph("ARQUIVOS CADASTRADOS", styles['Heading2']))
        if arquivos:
            arquivos_data = [['NB', 'APS', 'CPF Segurado', 'Tipo', 'Código Caixa']]
            for arquivo in arquivos:
                arquivos_data.append([
                    arquivo.get('NB', ''),
                    arquivo.get('APS', ''),
                    arquivo.get('SeguradoFK', ''),
                    arquivo.get('Tipo', ''),
                    arquivo.get('Caixa_codigo', '')
                ])
            arquivos_table = Table(arquivos_data, repeatRows=1)
            arquivos_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#27ae60')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#eafaf1')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(arquivos_table)
        else:
            elements.append(Paragraph("Nenhum arquivo cadastrado.", styles['Normal']))
        elements.append(Spacer(1, 20))

        # Segurados
        elements.append(Paragraph("SEGURADOS CADASTRADOS", styles['Heading2']))
        if segurados:
            segurados_data = [['CPF', 'Nome', 'Quantidade de Arquivos']]
            for segurado in segurados:
                num_arquivos = len(segurado.get('Arquivos', [])) if isinstance(segurado.get('Arquivos'), list) else 0
                segurados_data.append([
                    segurado.get('cpf', ''),
                    segurado.get('Nome', ''),
                    num_arquivos
                ])
            segurados_table = Table(segurados_data, repeatRows=1)
            segurados_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f39c12')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#fef5e7')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(segurados_table)
        else:
            elements.append(Paragraph("Nenhum segurado cadastrado.", styles['Normal']))

        # Rodapé
        elements.append(Spacer(1, 30))
        rodape = Paragraph(f"Relatório gerado automaticamente pelo Sistema GCA em {datetime.now().strftime('%d/%m/%Y às %H:%M')}", styles['Normal'])
        elements.append(rodape)

        # Gerar PDF
        doc.build(elements)
        return salvar_como

    def _gerar_excel_completo(self, estatisticas, caixas, arquivos, segurados, salvar_como):
        """Gera Excel do relatório completo"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Completo_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        wb = Workbook()
        
        # Estatísticas
        ws_stats = wb.active
        ws_stats.title = "Estatísticas"
        
        # Título
        ws_stats.merge_cells('A1:B1')
        ws_stats['A1'] = "RELATÓRIO COMPLETO - SISTEMA GCA"
        ws_stats['A1'].font = Font(size=16, bold=True, color='FF2C3E50')
        ws_stats['A1'].alignment = Alignment(horizontal='center')
        
        ws_stats['A3'] = "Data de geração:"
        ws_stats['B3'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        
        # Dados estatísticos
        ws_stats['A5'] = "ESTATÍSTICAS DO SISTEMA"
        ws_stats['A5'].font = Font(size=14, bold=True, color='FF2C3E50')
        
        headers = ['Descrição', 'Valor']
        ws_stats.append([''] * 2)  # Linha vazia
        ws_stats.append(headers)
        
        # Formatar cabeçalho
        for cell in ws_stats[7]:
            cell.font = Font(bold=True, color='FFFFFFFF')
            cell.fill = PatternFill(start_color='FF3498DB', end_color='FF3498DB', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            dados_stats = [
                ['Caixas Cadastradas', stats.get('caixas', 0)],
                ['Arquivos Cadastrados', stats.get('arquivos', 0)],
                ['Segurados Cadastrados', stats.get('segurados', 0)],
                ['Uptime do Sistema', estatisticas.get('uptime', 'N/A')]
            ]
            
            for dado in dados_stats:
                ws_stats.append(dado)
        
        # Caixas
        ws_caixas = wb.create_sheet("Caixas")
        ws_caixas.append(["RELATÓRIO DE CAIXAS"])
        ws_caixas.append([f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"])
        ws_caixas.append([])
        
        headers_caixas = ['Código', 'NBI', 'NBF', 'Prateleira', 'Bloco', 'Andar', 'Corredor']
        ws_caixas.append(headers_caixas)
        
        # Formatar cabeçalho caixas
        for cell in ws_caixas[4]:
            cell.font = Font(bold=True, color='FFFFFFFF')
            cell.fill = PatternFill(start_color='FFE74C3C', end_color='FFE74C3C', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        for caixa in caixas:
            ws_caixas.append([
                caixa.get('codigo', ''),
                caixa.get('NBI', ''),
                caixa.get('NBF', ''),
                caixa.get('prateleira', ''),
                caixa.get('bloco', ''),
                caixa.get('andar', ''),
                caixa.get('corredor', '')
            ])
        
        # Ajustar largura das colunas
        for column in ws_caixas.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws_caixas.column_dimensions[column_letter].width = adjusted_width
        
        # Arquivos
        ws_arquivos = wb.create_sheet("Arquivos")
        ws_arquivos.append(["RELATÓRIO DE ARQUIVOS"])
        ws_arquivos.append([f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"])
        ws_arquivos.append([])
        
        headers_arquivos = ['NB', 'APS', 'CPF Segurado', 'Tipo', 'Código Caixa']
        ws_arquivos.append(headers_arquivos)
        
        # Formatar cabeçalho arquivos
        for cell in ws_arquivos[4]:
            cell.font = Font(bold=True, color='FFFFFFFF')
            cell.fill = PatternFill(start_color='FF27AE60', end_color='FF27AE60', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        for arquivo in arquivos:
            ws_arquivos.append([
                arquivo.get('NB', ''),
                arquivo.get('APS', ''),
                arquivo.get('SeguradoFK', ''),
                arquivo.get('Tipo', ''),
                arquivo.get('Caixa_codigo', '')
            ])
        
        # Ajustar largura das colunas
        for column in ws_arquivos.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws_arquivos.column_dimensions[column_letter].width = adjusted_width
        
        # Segurados
        ws_segurados = wb.create_sheet("Segurados")
        ws_segurados.append(["RELATÓRIO DE SEGURADOS"])
        ws_segurados.append([f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"])
        ws_segurados.append([])
        
        headers_segurados = ['CPF', 'Nome', 'Quantidade de Arquivos']
        ws_segurados.append(headers_segurados)
        
        # Formatar cabeçalho segurados
        for cell in ws_segurados[4]:
            cell.font = Font(bold=True, color='FFFFFFFF')
            cell.fill = PatternFill(start_color='FFF39C12', end_color='FFF39C12', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        for segurado in segurados:
            num_arquivos = len(segurado.get('Arquivos', [])) if isinstance(segurado.get('Arquivos'), list) else 0
            ws_segurados.append([
                segurado.get('cpf', ''),
                segurado.get('Nome', ''),
                num_arquivos
            ])
        
        # Ajustar largura das colunas
        for column in ws_segurados.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws_segurados.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(salvar_como)
        return salvar_como

    def _gerar_pdf_caixas(self, caixas, estatisticas, salvar_como):
        """Gera PDF do relatório de caixas"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Caixas_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

        doc = SimpleDocTemplate(salvar_como, pagesize=A4)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            alignment=1,
            textColor=colors.HexColor('#2c3e50')
        )

        # Título
        title = Paragraph("RELATÓRIO DE CAIXAS - SISTEMA GCA", title_style)
        elements.append(title)

        # Data de geração
        data_geracao = Paragraph(f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", styles['Normal'])
        elements.append(data_geracao)
        elements.append(Spacer(1, 20))

        # Estatísticas
        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            elements.append(Paragraph(f"Total de Caixas: {stats.get('caixas', 0)}", styles['Heading3']))
            elements.append(Spacer(1, 10))

        # Tabela de caixas
        if caixas:
            caixas_data = [['Código', 'NBI', 'NBF', 'Prateleira', 'Bloco', 'Andar', 'Corredor']]
            for caixa in caixas:
                caixas_data.append([
                    caixa.get('codigo', ''),
                    caixa.get('NBI', ''),
                    caixa.get('NBF', ''),
                    caixa.get('prateleira', ''),
                    caixa.get('bloco', ''),
                    caixa.get('andar', ''),
                    caixa.get('corredor', '')
                ])
            
            caixas_table = Table(caixas_data, repeatRows=1)
            caixas_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e74c3c')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f9ebea')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(caixas_table)
        else:
            elements.append(Paragraph("Nenhuma caixa cadastrada.", styles['Normal']))

        elements.append(Spacer(1, 20))
        rodape = Paragraph(f"Relatório gerado automaticamente pelo Sistema GCA", styles['Normal'])
        elements.append(rodape)

        doc.build(elements)
        return salvar_como

    def _gerar_excel_caixas(self, caixas, estatisticas, salvar_como):
        """Gera Excel do relatório de caixas"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Caixas_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Caixas"
        
        # Cabeçalho
        ws.merge_cells('A1:G1')
        ws['A1'] = "RELATÓRIO DE CAIXAS - SISTEMA GCA"
        ws['A1'].font = Font(size=16, bold=True, color='FF2C3E50')
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws['A3'] = "Data de geração:"
        ws['B3'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        
        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            ws['A5'] = f"Total de Caixas: {stats.get('caixas', 0)}"
            ws['A5'].font = Font(size=12, bold=True)
        
        # Cabeçalho da tabela
        headers = ['Código', 'NBI', 'NBF', 'Prateleira', 'Bloco', 'Andar', 'Corredor']
        ws.append([])
        ws.append(headers)
        
        # Formatar cabeçalho
        for cell in ws[7]:
            cell.font = Font(bold=True, color='FFFFFFFF')
            cell.fill = PatternFill(start_color='FFE74C3C', end_color='FFE74C3C', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # Dados
        for caixa in caixas:
            ws.append([
                caixa.get('codigo', ''),
                caixa.get('NBI', ''),
                caixa.get('NBF', ''),
                caixa.get('prateleira', ''),
                caixa.get('bloco', ''),
                caixa.get('andar', ''),
                caixa.get('corredor', '')
            ])
        
        # Ajustar largura das colunas
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(salvar_como)
        return salvar_como

    def _gerar_pdf_arquivos(self, arquivos, estatisticas, salvar_como):
        """Gera PDF do relatório de arquivos"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Arquivos_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

        doc = SimpleDocTemplate(salvar_como, pagesize=A4)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            alignment=1,
            textColor=colors.HexColor('#2c3e50')
        )

        title = Paragraph("RELATÓRIO DE ARQUIVOS - SISTEMA GCA", title_style)
        elements.append(title)

        data_geracao = Paragraph(f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", styles['Normal'])
        elements.append(data_geracao)
        elements.append(Spacer(1, 20))

        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            elements.append(Paragraph(f"Total de Arquivos: {stats.get('arquivos', 0)}", styles['Heading3']))
            elements.append(Spacer(1, 10))

        if arquivos:
            arquivos_data = [['NB', 'APS', 'CPF Segurado', 'Tipo', 'Código Caixa']]
            for arquivo in arquivos:
                arquivos_data.append([
                    arquivo.get('NB', ''),
                    arquivo.get('APS', ''),
                    arquivo.get('SeguradoFK', ''),
                    arquivo.get('Tipo', ''),
                    arquivo.get('Caixa_codigo', '')
                ])
            
            arquivos_table = Table(arquivos_data, repeatRows=1)
            arquivos_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#27ae60')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#eafaf1')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(arquivos_table)
        else:
            elements.append(Paragraph("Nenhum arquivo cadastrado.", styles['Normal']))

        elements.append(Spacer(1, 20))
        rodape = Paragraph(f"Relatório gerado automaticamente pelo Sistema GCA", styles['Normal'])
        elements.append(rodape)

        doc.build(elements)
        return salvar_como

    def _gerar_excel_arquivos(self, arquivos, estatisticas, salvar_como):
        """Gera Excel do relatório de arquivos"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Arquivos_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Arquivos"
        
        ws.merge_cells('A1:E1')
        ws['A1'] = "RELATÓRIO DE ARQUIVOS - SISTEMA GCA"
        ws['A1'].font = Font(size=16, bold=True, color='FF2C3E50')
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws['A3'] = "Data de geração:"
        ws['B3'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        
        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            ws['A5'] = f"Total de Arquivos: {stats.get('arquivos', 0)}"
            ws['A5'].font = Font(size=12, bold=True)
        
        headers = ['NB', 'APS', 'CPF Segurado', 'Tipo', 'Código Caixa']
        ws.append([])
        ws.append(headers)
        
        for cell in ws[7]:
            cell.font = Font(bold=True, color='FFFFFFFF')
            cell.fill = PatternFill(start_color='FF27AE60', end_color='FF27AE60', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        for arquivo in arquivos:
            ws.append([
                arquivo.get('NB', ''),
                arquivo.get('APS', ''),
                arquivo.get('SeguradoFK', ''),
                arquivo.get('Tipo', ''),
                arquivo.get('Caixa_codigo', '')
            ])
        
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(salvar_como)
        return salvar_como

    def _gerar_pdf_segurados(self, segurados, estatisticas, salvar_como):
        """Gera PDF do relatório de segurados"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Segurados_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

        doc = SimpleDocTemplate(salvar_como, pagesize=A4)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            alignment=1,
            textColor=colors.HexColor('#2c3e50')
        )

        title = Paragraph("RELATÓRIO DE SEGURADOS - SISTEMA GCA", title_style)
        elements.append(title)

        data_geracao = Paragraph(f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", styles['Normal'])
        elements.append(data_geracao)
        elements.append(Spacer(1, 20))

        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            elements.append(Paragraph(f"Total de Segurados: {stats.get('segurados', 0)}", styles['Heading3']))
            elements.append(Spacer(1, 10))

        if segurados:
            segurados_data = [['CPF', 'Nome', 'Quantidade de Arquivos']]
            for segurado in segurados:
                num_arquivos = len(segurado.get('Arquivos', [])) if isinstance(segurado.get('Arquivos'), list) else 0
                segurados_data.append([
                    segurado.get('cpf', ''),
                    segurado.get('Nome', ''),
                    num_arquivos
                ])
            
            segurados_table = Table(segurados_data, repeatRows=1)
            segurados_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f39c12')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#fef5e7')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(segurados_table)
        else:
            elements.append(Paragraph("Nenhum segurado cadastrado.", styles['Normal']))

        elements.append(Spacer(1, 20))
        rodape = Paragraph(f"Relatório gerado automaticamente pelo Sistema GCA", styles['Normal'])
        elements.append(rodape)

        doc.build(elements)
        return salvar_como

    def _gerar_excel_segurados(self, segurados, estatisticas, salvar_como):
        """Gera Excel do relatório de segurados"""
        if salvar_como is None:
            salvar_como = f"Relatorio_Segurados_GCA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Segurados"
        
        ws.merge_cells('A1:C1')
        ws['A1'] = "RELATÓRIO DE SEGURADOS - SISTEMA GCA"
        ws['A1'].font = Font(size=16, bold=True, color='FF2C3E50')
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws['A3'] = "Data de geração:"
        ws['B3'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        
        if estatisticas:
            stats = estatisticas.get('estatisticas', {})
            ws['A5'] = f"Total de Segurados: {stats.get('segurados', 0)}"
            ws['A5'].font = Font(size=12, bold=True)
        
        headers = ['CPF', 'Nome', 'Quantidade de Arquivos']
        ws.append([])
        ws.append(headers)
        
        for cell in ws[7]:
            cell.font = Font(bold=True, color='FFFFFFFF')
            cell.fill = PatternFill(start_color='FFF39C12', end_color='FFF39C12', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        for segurado in segurados:
            num_arquivos = len(segurado.get('Arquivos', [])) if isinstance(segurado.get('Arquivos'), list) else 0
            ws.append([
                segurado.get('cpf', ''),
                segurado.get('Nome', ''),
                num_arquivos
            ])
        
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(salvar_como)
        return salvar_como

    def gerar_relatorio_personalizado(self, tipo_relatorio, filtros=None, formato='pdf', salvar_como=None):
        """Gera relatório personalizado com filtros específicos"""
        if filtros is None:
            filtros = {}
        
        # Implementar lógica de filtros personalizados
        # Esta é uma estrutura básica que pode ser expandida
        dados_filtrados = self._aplicar_filtros(tipo_relatorio, filtros)
        
        if formato.lower() == 'pdf':
            return self._gerar_pdf_personalizado(tipo_relatorio, dados_filtrados, salvar_como)
        elif formato.lower() == 'excel':
            return self._gerar_excel_personalizado(tipo_relatorio, dados_filtrados, salvar_como)
        else:
            raise ValueError("Formato deve ser 'pdf' ou 'excel'")

    def _aplicar_filtros(self, tipo_relatorio, filtros):
        """Aplica filtros aos dados (implementação básica)"""
        # Esta função pode ser expandida para aplicar filtros mais complexos
        dados = self.obter_dados_api(tipo_relatorio) or []
        
        # Exemplo de filtro simples por código
        if 'codigo' in filtros and filtros['codigo']:
            if tipo_relatorio == 'caixas':
                dados = [d for d in dados if str(d.get('codigo', '')).find(str(filtros['codigo'])) != -1]
            elif tipo_relatorio == 'arquivos':
                dados = [d for d in dados if str(d.get('NB', '')).find(str(filtros['codigo'])) != -1]
            elif tipo_relatorio == 'segurados':
                dados = [d for d in dados if str(d.get('cpf', '')).find(str(filtros['codigo'])) != -1]
        
        return dados

    def _gerar_pdf_personalizado(self, tipo_relatorio, dados, salvar_como):
        """Gera PDF personalizado (estrutura básica)"""
        # Implementação similar às outras funções PDF, mas personalizável
        pass

    def _gerar_excel_personalizado(self, tipo_relatorio, dados, salvar_como):
        """Gera Excel personalizado (estrutura básica)"""
        # Implementação similar às outras funções Excel, mas personalizável
        pass


# Interface para teste do módulo
class TesteRelatorios:
    def __init__(self):
        self.gerador = GeradorRelatorios()
        
    def testar_conexao(self):
        if self.gerador.testar_conexao_api():
            print("✅ Conexão com a API estabelecida com sucesso!")
            return True
        else:
            print("❌ Não foi possível conectar à API.")
            return False
    
    def menu_principal(self):
        while True:
            print("\n" + "="*50)
            print("SISTEMA DE RELATÓRIOS GCA")
            print("="*50)
            print("1. Relatório Completo (PDF)")
            print("2. Relatório Completo (Excel)")
            print("3. Relatório de Caixas (PDF)")
            print("4. Relatório de Caixas (Excel)")
            print("5. Relatório de Arquivos (PDF)")
            print("6. Relatório de Arquivos (Excel)")
            print("7. Relatório de Segurados (PDF)")
            print("8. Relatório de Segurados (Excel)")
            print("9. Testar Conexão com API")
            print("0. Sair")
            
            opcao = input("\nEscolha uma opção: ")
            
            try:
                if opcao == '1':
                    caminho = self.gerador.gerar_relatorio_completo('pdf')
                    print(f"✅ Relatório completo PDF gerado: {caminho}")
                elif opcao == '2':
                    caminho = self.gerador.gerar_relatorio_completo('excel')
                    print(f"✅ Relatório completo Excel gerado: {caminho}")
                elif opcao == '3':
                    caminho = self.gerador.gerar_relatorio_caixas('pdf')
                    print(f"✅ Relatório de caixas PDF gerado: {caminho}")
                elif opcao == '4':
                    caminho = self.gerador.gerar_relatorio_caixas('excel')
                    print(f"✅ Relatório de caixas Excel gerado: {caminho}")
                elif opcao == '5':
                    caminho = self.gerador.gerar_relatorio_arquivos('pdf')
                    print(f"✅ Relatório de arquivos PDF gerado: {caminho}")
                elif opcao == '6':
                    caminho = self.gerador.gerar_relatorio_arquivos('excel')
                    print(f"✅ Relatório de arquivos Excel gerado: {caminho}")
                elif opcao == '7':
                    caminho = self.gerador.gerar_relatorio_segurados('pdf')
                    print(f"✅ Relatório de segurados PDF gerado: {caminho}")
                elif opcao == '8':
                    caminho = self.gerador.gerar_relatorio_segurados('excel')
                    print(f"✅ Relatório de segurados Excel gerado: {caminho}")
                elif opcao == '9':
                    self.testar_conexao()
                elif opcao == '0':
                    print("Saindo...")
                    break
                else:
                    print("❌ Opção inválida!")
            except Exception as e:
                print(f"❌ Erro ao gerar relatório: {e}")


if __name__ == "__main__":
    teste = TesteRelatorios()
    
    if teste.testar_conexao():
        teste.menu_principal()
    else:
        print("Não foi possível iniciar o sistema de relatórios.")